"""Extract text from every document in data/work/inventory.csv into data/work/text/.

One Markdown file per document, named <display_title slug>-<sha256[:8]>.md (never the original
file name, which can contain people's names). Each file starts with a YAML-style header:
path, display_title, track, tracks, kind, file_type, sha256, method, ocr_pages, ocr_note.
The path field is for tracing inside data/work/ only: it must never be loaded into Supabase or
shown in the app.

  - PDF:  "## Page N" per page. Pages with under MIN_CHARS_PER_PAGE selectable characters are
          OCR'd (300 dpi, Tesseract eng). For kind=reference only the first OCR_LIMIT_REFERENCE
          such pages are OCR'd. Link URIs are listed under each page.
  - DOCX: headings, lists and tables in document order, hyperlinks inline as [text](url).
  - PPTX: "## Slide N", hyperlinks inline, speaker notes.
  - MD / TXT: copied as-is.
  - .doc / .ppt: converted with LibreOffice (soffice) if installed, else logged as a failure.
  - kind=skip rows (e.g. code files) are not extracted.

Resumable: a document whose output exists with the same sha256 is skipped (--force redoes all).
Failures go to data/work/extract_errors.csv; one bad file never stops the run.

Run from the repo root:  python scripts/extract.py [--force] [--only SUBSTRING]
"""

import argparse
import csv
import re
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import datetime
from pathlib import Path

import pymupdf
import pytesseract
from docx import Document
from docx.text.hyperlink import Hyperlink
from PIL import Image
from pptx import Presentation

from inventory import MIN_CHARS_PER_PAGE, RAW_DIR, ROOT, TEXT_SUFFIXES

INVENTORY = ROOT / "data" / "work" / "inventory.csv"
OUT_DIR = ROOT / "data" / "work" / "text"
ERRORS_CSV = ROOT / "data" / "work" / "extract_errors.csv"
CONVERT_DIR = ROOT / "data" / "work" / "convert"

OCR_DPI = 300
OCR_LANG = "eng"
OCR_LIMIT_REFERENCE = 15  # reference files only need a summary and topics
PROGRESS_EVERY = 25  # print a line every N OCR'd pages
LIBREOFFICE_TARGET = {".doc": "docx", ".ppt": "pptx"}

HEADER_FIELDS = ["path", "display_title", "track", "tracks", "kind", "file_type", "sha256"]


class Failures:
    """Collects failures for extract_errors.csv."""

    def __init__(self):
        self.rows = []

    def add(self, row, stage, error):
        msg = f"{type(error).__name__}: {error}" if isinstance(error, Exception) else str(error)
        print(f"  FAIL [{stage}] {msg}", file=sys.stderr)
        self.rows.append({
            "path": row["path"],
            "display_title": row["display_title"],
            "stage": stage,
            "error": msg,
            "time": datetime.now().isoformat(timespec="seconds"),
        })

    def write(self):
        with ERRORS_CSV.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=["path", "display_title", "stage", "error", "time"])
            writer.writeheader()
            writer.writerows(self.rows)


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60] or "document"


def out_path(row):
    return OUT_DIR / f"{slug(row['display_title'])}-{row['sha256'][:8]}.md"


def is_up_to_date(path, sha):
    if not path.exists():
        return False
    with path.open(encoding="utf-8") as f:
        next(f, None)  # opening ---
        for line in f:
            if line.startswith("sha256:"):
                return line.split(":", 1)[1].strip() == sha
            if line.strip() == "---":  # end of header
                break
    return False


def clean(text):
    """Trim trailing spaces and collapse runs of blank lines."""
    text = "\n".join(line.rstrip() for line in text.splitlines())
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ---------- PDF ----------

def ocr_page(page):
    pix = page.get_pixmap(dpi=OCR_DPI)
    img = Image.frombytes("RGB" if pix.n < 4 else "RGBA", (pix.width, pix.height), pix.samples)
    return pytesseract.image_to_string(img, lang=OCR_LANG)


def page_links(page):
    lines = []
    seen = set()
    for link in page.get_links():
        uri = link.get("uri")
        if not uri or uri in seen:
            continue
        seen.add(uri)
        anchor = " ".join(page.get_textbox(link["from"]).split())
        lines.append(f"- [{anchor}]({uri})" if anchor and anchor != uri else f"- {uri}")
    return lines


def extract_pdf(path, row, failures):
    limit = OCR_LIMIT_REFERENCE if row["kind"] == "reference" else None
    parts, ocr_pages, not_ocred = [], [], []
    with pymupdf.open(path) as doc:
        for page in doc:
            n = page.number + 1
            text = page.get_text()
            if len(text.strip()) < MIN_CHARS_PER_PAGE:
                if limit is not None and len(ocr_pages) >= limit:
                    not_ocred.append(n)
                    text = (text + "\n[not OCR'd]").strip()
                else:
                    try:
                        text = ocr_page(page)
                        ocr_pages.append(n)
                        if len(ocr_pages) % PROGRESS_EVERY == 0:
                            print(f"  ... {len(ocr_pages)} pages OCR'd (page {n}/{doc.page_count})")
                    except Exception as e:
                        failures.add(row, f"ocr page {n}", e)
                        text = (text + "\n[OCR failed]").strip()
            body = clean(text)
            links = page_links(page)
            if links:
                body += "\n\nLinks:\n" + "\n".join(links)
            parts.append(f"## Page {n}\n\n{body}")
        total = doc.page_count

    if not ocr_pages:
        method = "text"
    elif len(ocr_pages) == total:
        method = "ocr"
    else:
        method = "mixed"
    meta = {"method": method, "ocr_pages": ", ".join(map(str, ocr_pages))}
    if not_ocred:
        meta["ocr_note"] = f"OCR limited to {OCR_LIMIT_REFERENCE} pages (reference)"
    return "\n\n".join(parts), meta


# ---------- DOCX ----------

def docx_paragraph(p):
    text = ""
    for item in p.iter_inner_content():
        if isinstance(item, Hyperlink):
            label = item.text.strip()
            url = item.url
            # Word sometimes splits one link into pieces whose labels are URL fragments
            if url and label and label not in url:
                text += f"[{label}]({url})"
            else:
                text += url or label
        else:
            text += item.text
    text = text.strip()
    if not text:
        return ""
    style = (p.style.name if p.style is not None else "").lower()
    m = re.match(r"heading (\d)", style)
    if m:
        return "#" * min(int(m.group(1)) + 1, 6) + " " + text
    if style == "title":
        return "# " + text
    if "list" in style or p._p.pPr is not None and p._p.pPr.numPr is not None:
        return "- " + text
    return text


def docx_table(table):
    lines = []
    for r in table.rows:
        cells = []
        for c in r.cells:
            cell = " ".join(docx_paragraph(p) for p in c.paragraphs).strip()
            if not cells or cell != cells[-1]:  # merged cells repeat their text
                cells.append(cell)
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_docx(path):
    doc = Document(path)
    blocks = []
    for item in doc.iter_inner_content():
        block = docx_table(item) if hasattr(item, "rows") else docx_paragraph(item)
        if block:
            blocks.append(block)
    return clean("\n\n".join(blocks)), {"method": "text"}


# ---------- PPTX ----------

def pptx_text_frame(frame):
    lines = []
    for para in frame.paragraphs:
        line = ""
        for run in para.runs:
            url = run.hyperlink.address if run.hyperlink is not None else None
            line += f"[{run.text}]({url})" if url and run.text.strip() else run.text
        if line.strip():
            lines.append(line.strip())
    return "\n".join(lines)


def extract_pptx(path):
    prs = Presentation(path)
    parts = []
    for n, slide in enumerate(prs.slides, 1):
        texts = [pptx_text_frame(s.text_frame) for s in slide.shapes if s.has_text_frame]
        body = "\n\n".join(t for t in texts if t)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                body += f"\n\nNotes:\n{notes}"
        parts.append(f"## Slide {n}\n\n{clean(body)}")
    return "\n\n".join(parts), {"method": "text"}


# ---------- old formats ----------

def extract_via_libreoffice(path, row, failures):
    if not shutil.which("soffice"):
        failures.add(row, "convert", "LibreOffice not installed (sudo apt install libreoffice)")
        return None
    target = LIBREOFFICE_TARGET[path.suffix.lower()]
    CONVERT_DIR.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=CONVERT_DIR) as tmp:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", target, "--outdir", tmp, str(path)],
            check=True, capture_output=True, timeout=300,
        )
        converted = Path(tmp) / f"{path.stem}.{target}"
        body, meta = extract_docx(converted) if target == "docx" else extract_pptx(converted)
    meta["method"] = f"libreoffice+{target}"
    return body, meta


def extract(path, row, failures):
    """Returns (markdown body, header extras) or None on failure."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path, row, failures)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix in LIBREOFFICE_TARGET:
        return extract_via_libreoffice(path, row, failures)
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="ignore").strip(), {"method": "text"}
    failures.add(row, "extract", f"unsupported type {suffix or '(none)'}")
    return None


def render(row, body, meta):
    header = {k: row[k] for k in HEADER_FIELDS}
    header.update({"method": meta["method"], "ocr_pages": meta.get("ocr_pages", "")})
    if meta.get("ocr_note"):
        header["ocr_note"] = meta["ocr_note"]
    lines = ["---"] + [f"{k}: {v}" for k, v in header.items()] + ["---", ""]
    return "\n".join(lines) + f"# {row['display_title']}\n\n{body}\n"


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--force", action="store_true", help="re-extract even if up to date")
    parser.add_argument("--only", help="only rows whose path contains this text")
    args = parser.parse_args()

    if not INVENTORY.exists():
        sys.exit(f"{INVENTORY} not found: run scripts/inventory.py first")
    with INVENTORY.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    if args.only:
        rows = [r for r in rows if args.only in r["path"]]

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    failures = Failures()
    done = skipped = excluded = failed = ocr_total = 0
    start = time.time()

    for i, row in enumerate(rows, 1):
        title = row["display_title"]
        if row["kind"] == "skip":
            excluded += 1
            continue
        out = out_path(row)
        if not args.force and is_up_to_date(out, row["sha256"]):
            skipped += 1
            continue
        print(f"[{i}/{len(rows)}] {title}")
        t0 = time.time()
        try:
            result = extract(RAW_DIR / row["path"], row, failures)
        except Exception as e:  # corrupt/encrypted file: record it, keep going
            failures.add(row, "extract", e)
            result = None
        if result is None:
            failed += 1
            continue
        body, meta = result
        tmp = out.with_suffix(".md.tmp")
        tmp.write_text(render(row, body, meta), encoding="utf-8")
        tmp.replace(out)
        n_ocr = len(meta["ocr_pages"].split(", ")) if meta.get("ocr_pages") else 0
        ocr_total += n_ocr
        note = f", {n_ocr} pages OCR'd" if n_ocr else ""
        print(f"  -> {out.relative_to(ROOT)} ({meta['method']}{note}, {time.time() - t0:.0f}s)")
        done += 1

    failures.write()
    print(f"\nDone {done}, up to date {skipped}, kind=skip {excluded}, failed {failed}; "
          f"{ocr_total} pages OCR'd; {len(failures.rows)} errors in {ERRORS_CSV.relative_to(ROOT)}; "
          f"{time.time() - start:.0f}s")


if __name__ == "__main__":
    main()
