"""Inventory every file under data/raw and write data/work/inventory.csv.

Columns: path, track, subfolder, file_type, size_bytes, pages, text_status
  - track: top-level folder under data/raw (e.g. drive, cdc, curriculum, notion)
  - subfolder: folders between the track and the file (drive keeps its per-track folders here)
  - pages: PDF pages, PPTX slides, DOCX pages (from document properties, may be blank)
  - text_status: text | ocr | mixed | unknown
      text  = selectable text found, no OCR needed
      ocr   = no (or almost no) selectable text, needs OCR
      mixed = PDF where only some pages have selectable text (those pages need OCR)

Run from the repo root:  python scripts/inventory.py
"""

import csv
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = ROOT / "data" / "raw"
OUT_CSV = ROOT / "data" / "work" / "inventory.csv"

# A PDF page with fewer characters than this is treated as having no real text layer.
MIN_CHARS_PER_PAGE = 25
TEXT_SUFFIXES = {".md", ".txt", ".cpp", ".py", ".csv", ".json", ".html"}
SKIP_DIRS = {"excluded"}


def inspect_pdf(path):
    with pymupdf.open(path) as doc:
        pages = doc.page_count
        with_text = sum(
            1 for page in doc if len(page.get_text().strip()) >= MIN_CHARS_PER_PAGE
        )
    if pages == 0 or with_text == 0:
        status = "ocr"
    elif with_text == pages:
        status = "text"
    else:
        status = "mixed"
    return pages, status


def inspect_docx(path):
    doc = Document(path)
    chars = sum(len(p.text.strip()) for p in doc.paragraphs)
    chars += sum(
        len(cell.text.strip())
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    pages = ""
    with zipfile.ZipFile(path) as z:
        if "docProps/app.xml" in z.namelist():
            m = re.search(r"<Pages>(\d+)</Pages>", z.read("docProps/app.xml").decode("utf-8", "ignore"))
            if m:
                pages = int(m.group(1))
    return pages, "text" if chars >= MIN_CHARS_PER_PAGE else "ocr"


def inspect_pptx(path):
    prs = Presentation(path)
    chars = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chars += len(shape.text_frame.text.strip())
    return len(prs.slides), "text" if chars >= MIN_CHARS_PER_PAGE else "ocr"


def inspect_plain(path):
    text = path.read_text(encoding="utf-8", errors="ignore")
    return "", "text" if text.strip() else "unknown"


def inspect(path):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return inspect_pdf(path)
    if suffix == ".docx":
        return inspect_docx(path)
    if suffix == ".pptx":
        return inspect_pptx(path)
    if suffix in TEXT_SUFFIXES:
        return inspect_plain(path)
    return "", "unknown"


def main():
    if not RAW_DIR.is_dir():
        sys.exit(f"{RAW_DIR} not found")

    files = sorted(
        p for p in RAW_DIR.rglob("*")
        if p.is_file() and not (set(p.relative_to(RAW_DIR).parts) & SKIP_DIRS)
    )

    rows = []
    for path in files:
        rel = path.relative_to(RAW_DIR)
        try:
            pages, status = inspect(path)
        except Exception as e:  # corrupt/encrypted file: record it, keep going
            print(f"WARN {rel}: {type(e).__name__}: {e}", file=sys.stderr)
            pages, status = "", "unknown"
        rows.append({
            "path": str(rel),
            "track": rel.parts[0] if len(rel.parts) > 1 else "(root)",
            "subfolder": "/".join(rel.parts[1:-1]),
            "file_type": path.suffix.lower().lstrip(".") or "none",
            "size_bytes": path.stat().st_size,
            "pages": pages,
            "text_status": status,
        })

    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else ["path"])
        writer.writeheader()
        writer.writerows(rows)

    print(f"Wrote {len(rows)} files to {OUT_CSV.relative_to(ROOT)}\n")
    print_summary(rows)


def print_summary(rows):
    by_track = defaultdict(Counter)
    size = Counter()
    pages = Counter()
    for r in rows:
        by_track[r["track"]][r["file_type"]] += 1
        size[r["track"]] += r["size_bytes"]
        pages[r["track"]] += r["pages"] or 0

    print(f"{'track':<14}{'type':<8}{'files':>6}")
    for track in sorted(by_track):
        for ftype, n in sorted(by_track[track].items()):
            print(f"{track:<14}{ftype:<8}{n:>6}")
        print(f"{'':<14}{'total':<8}{sum(by_track[track].values()):>6}"
              f"  {size[track] / 1e6:.1f} MB, {pages[track]} pages/slides")

    print("\nText extraction status:")
    status = Counter((r["file_type"], r["text_status"]) for r in rows)
    for (ftype, st), n in sorted(status.items()):
        print(f"  {ftype:<8}{st:<9}{n:>4}")

    needs_ocr = [r for r in rows if r["text_status"] in ("ocr", "mixed")]
    if needs_ocr:
        print(f"\nNeed OCR ({len(needs_ocr)}):")
        for r in needs_ocr:
            print(f"  [{r['text_status']}] {r['path']}")


if __name__ == "__main__":
    main()
